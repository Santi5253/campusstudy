import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def add_inline(paragraph, text):
    parts = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def _heading(doc, text, style):
    para = doc.add_paragraph(style=style)
    add_inline(para, text)
    return para


def build(md_text, out_path):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(6)
            run = para.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            header, body = rows[0], rows[2:]
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for c, text in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.text = ""
                add_inline(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in body:
                cells = table.add_row().cells
                for c, text in enumerate(row):
                    if c < len(cells):
                        cells[c].text = ""
                        add_inline(cells[c].paragraphs[0], text)
            doc.add_paragraph()
            continue

        if stripped == "---":
            doc.add_paragraph().add_run("_" * 60).font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
            i += 1
            continue

        if stripped.startswith("#### "):
            _heading(doc, stripped[5:], "Heading 3")
        elif stripped.startswith("### "):
            _heading(doc, stripped[4:], "Heading 2")
        elif stripped.startswith("## "):
            _heading(doc, stripped[3:], "Heading 1")
        elif stripped.startswith("# "):
            _heading(doc, stripped[2:], "Title")
        elif stripped.startswith("> "):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(18)
            add_inline(para, stripped[2:])
            for run in para.runs:
                run.italic = True
        elif re.match(r"^[-*] ", stripped):
            para = doc.add_paragraph(style="List Bullet")
            add_inline(para, stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            para = doc.add_paragraph(style="List Number")
            add_inline(para, re.sub(r"^\d+\. ", "", stripped))
        elif stripped == "":
            pass
        else:
            para = doc.add_paragraph()
            add_inline(para, line)
        i += 1

    doc.save(out_path)


if __name__ == "__main__":
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    build(src.read_text(encoding="utf-8"), dst)
    print("generado:", dst)
